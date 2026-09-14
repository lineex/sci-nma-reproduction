"""
Search Flow Conservation Audit Ledger Generator.
Tracks exact hit counts, batch export landing file completeness, and generates
publication-grade Search Audit Tables (Excel/Markdown) ensuring PRISMA Flow Conservation (L ≡ 0).
"""

import os
import json
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class SearchAuditLedger:
    """
    Constructs and audits the PRISMA Search Flow Conservation Ledger.
    """

    @classmethod
    def compile_audit_ledger(
        cls,
        search_queries: Dict[str, str],
        db_hit_counts: Dict[str, int],
        ingestion_metadata: Dict[str, Any],
        dedup_metrics: Dict[str, Any],
        search_date: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Compile comprehensive audit ledger comparing reported hits vs landed export records.
        """
        if not search_date:
            search_date = datetime.now().strftime("%Y-%m-%d")

        ledger = {
            "search_date": search_date,
            "databases": {},
            "total_reported_hits": 0,
            "total_landed_records": 0,
            "total_identified": 0,
            "duplicates_removed": dedup_metrics.get("duplicates_removed", 0),
            "records_to_screen": dedup_metrics.get("unique_records", 0),
            "conservation_status": "PENDING",
            "audit_discrepancies": []
        }

        db_names = ["PubMed", "Embase", "Web of Science", "Cochrane Library"]

        for db in db_names:
            reported = db_hit_counts.get(db, 0)
            ingested_info = ingestion_metadata.get("databases", {}).get(db, {})
            landed = ingested_info.get("record_count", 0)
            files = [f["filename"] for f in ingested_info.get("files", [])]

            query = search_queries.get(db, "")

            is_match = (reported == landed) if reported > 0 else True
            if not is_match:
                ledger["audit_discrepancies"].append(
                    f"[{db}] Reported hit count ({reported}) does not equal landed export count ({landed})."
                )

            ledger["databases"][db] = {
                "reported_hits": reported,
                "landed_records": landed,
                "batch_files": files,
                "query": query,
                "is_match": is_match
            }
            ledger["total_reported_hits"] += reported
            ledger["total_landed_records"] += landed

        # Set total identified from landed records
        ledger["total_identified"] = ledger["total_landed_records"]

        # Validate flow conservation
        # L = total_identified - duplicates_removed - records_to_screen
        loss = ledger["total_identified"] - ledger["duplicates_removed"] - ledger["records_to_screen"]
        if loss == 0 and len(ledger["audit_discrepancies"]) == 0:
            ledger["conservation_status"] = "PASSED (L ≡ 0)"
        else:
            ledger["conservation_status"] = f"FAILED (Loss L = {loss})"

        return ledger

    @classmethod
    def export_excel_audit_table(cls, ledger: Dict[str, Any], output_path: str) -> str:
        """Export publication-grade Excel Search Audit Table."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Search Audit Ledger"
        ws.views.sheetView[0].showGridLines = True

        # Styles
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        sub_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        pass_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        warn_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
        title_font = Font(name="Arial", size=14, bold=True, color="1F4E78")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        bold_font = Font(name="Arial", size=10, bold=True)
        norm_font = Font(name="Arial", size=10)
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        # Title
        ws.merge_cells("A1:G1")
        ws["A1"] = f"PRISMA 2020 Multi-Database Search Flow Conservation & Audit Table ({ledger.get('search_date')})"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 30

        # Subtitle
        ws.merge_cells("A2:G2")
        ws["A2"] = f"Audit Status: {ledger.get('conservation_status')} | Total Landed: {ledger.get('total_identified')} | Duplicates: {ledger.get('duplicates_removed')} | Unique to Screen: {ledger.get('records_to_screen')}"
        ws["A2"].font = Font(name="Arial", size=10, italic=True)
        ws.row_dimensions[2].height = 20

        # Headers
        headers = [
            "Database / Information Source",
            "Search Date",
            "Boolean Search Syntax",
            "Database Hit Count",
            "Landed Export Files",
            "Landed Records",
            "Audit Check (Match)"
        ]
        ws.row_dimensions[4].height = 25
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Data rows
        curr_row = 5
        for db, info in ledger.get("databases", {}).items():
            ws.row_dimensions[curr_row].height = 22
            ws.cell(row=curr_row, column=1, value=db).font = bold_font
            ws.cell(row=curr_row, column=2, value=ledger.get("search_date")).font = norm_font
            ws.cell(row=curr_row, column=3, value=info.get("query", "")[:120] + "...").font = norm_font
            ws.cell(row=curr_row, column=4, value=info.get("reported_hits", 0)).font = norm_font
            ws.cell(row=curr_row, column=5, value=", ".join(info.get("batch_files", []))).font = norm_font
            ws.cell(row=curr_row, column=6, value=info.get("landed_records", 0)).font = norm_font

            match_cell = ws.cell(row=curr_row, column=7, value="PASS" if info.get("is_match") else "MISMATCH")
            match_cell.font = bold_font
            match_cell.alignment = Alignment(horizontal="center")
            match_cell.fill = pass_fill if info.get("is_match") else warn_fill

            for c in range(1, 8):
                ws.cell(row=curr_row, column=c).border = thin_border

            curr_row += 1

        # Summary Rows
        ws.row_dimensions[curr_row].height = 24
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=3)
        ws.cell(row=curr_row, column=1, value="TOTAL IDENTIFIED RECORDS (Σ Database Hits)").font = bold_font
        ws.cell(row=curr_row, column=1).fill = sub_header_fill
        ws.cell(row=curr_row, column=4, value=ledger.get("total_reported_hits", 0)).font = bold_font
        ws.cell(row=curr_row, column=4).fill = sub_header_fill
        ws.cell(row=curr_row, column=6, value=ledger.get("total_landed_records", 0)).font = bold_font
        ws.cell(row=curr_row, column=6).fill = sub_header_fill
        ws.cell(row=curr_row, column=7, value="CONSERVED" if ledger.get("conservation_status", "").startswith("PASSED") else "FAIL").font = bold_font
        ws.cell(row=curr_row, column=7).fill = pass_fill if ledger.get("conservation_status", "").startswith("PASSED") else warn_fill
        for c in range(1, 8):
            ws.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

        # Duplicates Row
        ws.row_dimensions[curr_row].height = 22
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=5)
        ws.cell(row=curr_row, column=1, value="Duplicates Removed (Multi-Tier Provenance Deduplication)").font = bold_font
        ws.cell(row=curr_row, column=6, value=ledger.get("duplicates_removed", 0)).font = bold_font
        for c in range(1, 8):
            ws.cell(row=curr_row, column=c).border = thin_border
        curr_row += 1

        # Screened Row
        ws.row_dimensions[curr_row].height = 22
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=5)
        ws.cell(row=curr_row, column=1, value="NET UNIQUE RECORDS ADVANCED TO TITLE/ABSTRACT SCREENING").font = bold_font
        ws.cell(row=curr_row, column=6, value=ledger.get("records_to_screen", 0)).font = bold_font
        ws.cell(row=curr_row, column=6).fill = pass_fill
        for c in range(1, 8):
            ws.cell(row=curr_row, column=c).border = thin_border

        # Adjust column widths
        ws.column_dimensions["A"].width = 25
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 40
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 30
        ws.column_dimensions["F"].width = 18
        ws.column_dimensions["G"].width = 20

        wb.save(output_path)
        return output_path

    @classmethod
    def generate_prisma_flow_json(cls, ledger: Dict[str, Any], output_path: str) -> str:
        """Create draft prisma_flow.json compatible with Gate 1."""
        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        databases = {db: info["landed_records"] for db, info in ledger.get("databases", {}).items()}

        flow = {
            "databases": databases,
            "registries_or_citations": 0,
            "total_identified": ledger.get("total_identified", 0),
            "duplicates_removed": ledger.get("duplicates_removed", 0),
            "records_screened": ledger.get("records_to_screen", 0),
            "screening_excluded": 0,
            "screening_exclusion_reasons": {},
            "reports_sought": ledger.get("records_to_screen", 0),
            "reports_not_retrieved": 0,
            "reports_assessed": ledger.get("records_to_screen", 0),
            "fulltext_excluded": 0,
            "fulltext_exclusion_reasons": {
                "Wrong Population": 0,
                "Wrong Intervention": 0,
                "No Control Group": 0,
                "Ineligible Study Design": 0,
                "Duplicate Cohort": 0
            },
            "studies_included": 0
        }
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(flow, f, indent=2)
        return output_path
