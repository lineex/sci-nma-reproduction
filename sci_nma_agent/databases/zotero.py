"""
Zotero full-text bridge for evidence-synthesis projects.

The bridge is deliberately read-only with respect to Zotero. It can inspect a
local Zotero data directory (``zotero.sqlite`` + ``storage/``) or a Zotero
JSON export, resolve parent items to their PDF/FT-cache attachments, and
materialize a provenance-preserving analysis corpus inside a review project.
"""

from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, Iterable, Iterator, List, Optional, Sequence, Tuple
from urllib.parse import unquote, urlparse


class ZoteroError(RuntimeError):
    """Raised when a Zotero source cannot be read or materialized."""


def _utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def _sha256(path: Path) -> Optional[str]:
    if not path.is_file():
        return None
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _normalise_doi(value: Any) -> Optional[str]:
    if not value:
        return None
    text = str(value).strip().lower()
    text = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", text)
    text = re.sub(r"^doi:\s*", "", text)
    match = re.search(r"10\.\d{4,9}/[-._;()/:a-z0-9]+", text, re.I)
    return match.group(0).rstrip(".,;") if match else text or None


def _normalise_pmid(value: Any) -> Optional[str]:
    if not value:
        return None
    match = re.search(r"\b\d{6,9}\b", str(value))
    return match.group(0) if match else str(value).strip() or None


def _year(value: Any) -> Optional[int]:
    match = re.search(r"\b(19|20)\d{2}\b", str(value or ""))
    return int(match.group(0)) if match else None


def _slug(value: str, fallback: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", value or "").strip("_")
    return (slug[:100] or fallback).strip("_")


def _safe_json(value: Any) -> Any:
    """Convert common SQLite/export values to JSON-safe primitives."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, dict):
        return {str(k): _safe_json(v) for k, v in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [_safe_json(v) for v in value]
    return str(value)


@dataclass
class ZoteroAttachment:
    key: str
    parent_key: str
    content_type: Optional[str] = None
    link_mode: Optional[int] = None
    zotero_path: Optional[str] = None
    source_path: Optional[str] = None
    fulltext_cache_path: Optional[str] = None

    @property
    def pdf_path(self) -> Optional[Path]:
        if self.source_path and self.source_path.lower().endswith(".pdf"):
            return Path(self.source_path)
        return None

    @property
    def has_local_content(self) -> bool:
        return bool(
            (self.source_path and Path(self.source_path).is_file())
            or (self.fulltext_cache_path and Path(self.fulltext_cache_path).is_file())
        )


@dataclass
class ZoteroItem:
    key: str
    title: str = ""
    item_type: str = "journalArticle"
    doi: Optional[str] = None
    pmid: Optional[str] = None
    year: Optional[int] = None
    journal: Optional[str] = None
    abstract: Optional[str] = None
    authors: List[str] = field(default_factory=list)
    extra: Optional[str] = None
    collection_keys: List[str] = field(default_factory=list)
    attachments: List[ZoteroAttachment] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        return _safe_json(asdict(self))


class ZoteroLibrary:
    """Read-only adapter for a Zotero data directory or JSON export."""

    def __init__(self, data_dir: Optional[str] = None, export_json: Optional[str] = None):
        if not data_dir and not export_json:
            raise ZoteroError("Provide a Zotero data directory or JSON export.")
        self.data_dir = Path(data_dir).expanduser().resolve() if data_dir else None
        self.export_json = Path(export_json).expanduser().resolve() if export_json else None
        self.source_kind = "sqlite" if self.data_dir else "json"

    @classmethod
    def from_source(cls, source: str) -> "ZoteroLibrary":
        path = Path(source).expanduser().resolve()
        if path.is_file() and path.suffix.lower() == ".json":
            return cls(export_json=str(path))
        if path.is_dir():
            return cls(data_dir=str(path))
        raise ZoteroError(f"Zotero source not found: {source}")

    def load_items(self, collection: Optional[str] = None) -> List[ZoteroItem]:
        items = self._load_json() if self.export_json else self._load_sqlite()
        if collection:
            needle = collection.casefold()
            items = [
                item for item in items
                if needle in {key.casefold() for key in item.collection_keys}
                or needle == item.key.casefold()
            ]
        return items

    # ------------------------------ JSON export ---------------------------
    def _load_json(self) -> List[ZoteroItem]:
        assert self.export_json is not None
        try:
            raw = json.loads(self.export_json.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError) as exc:
            raise ZoteroError(f"Unable to read Zotero JSON export: {self.export_json}: {exc}") from exc

        records = raw.get("items", raw) if isinstance(raw, dict) else raw
        if not isinstance(records, list):
            raise ZoteroError("Zotero JSON export must contain an items list.")

        parent_by_key: Dict[str, ZoteroItem] = {}
        attachment_records: List[Dict[str, Any]] = []
        for record in records:
            if not isinstance(record, dict):
                continue
            data = record.get("data") if isinstance(record.get("data"), dict) else record
            item_type = str(data.get("itemType", record.get("itemType", "")) or "")
            key = str(data.get("key", record.get("key", "")) or "").strip()
            if not key:
                continue
            if item_type.lower() == "attachment" or record.get("parentItem"):
                attachment_records.append(record)
                continue

            item = self._item_from_export(data, key)
            parent_by_key[key] = item
            for nested in data.get("attachments", record.get("attachments", [])) or []:
                if isinstance(nested, dict):
                    nested = dict(nested)
                    nested.setdefault("parentItem", key)
                    attachment_records.append(nested)

        for record in attachment_records:
            data = record.get("data") if isinstance(record.get("data"), dict) else record
            parent_key = str(
                data.get("parentItem", record.get("parentItem", "")) or ""
            ).strip()
            parent = parent_by_key.get(parent_key)
            if not parent:
                continue
            attachment = self._attachment_from_export(data, parent.key)
            if attachment:
                parent.attachments.append(attachment)

        return list(parent_by_key.values())

    def _item_from_export(self, data: Dict[str, Any], key: str) -> ZoteroItem:
        authors: List[str] = []
        for author in data.get("creators", []) or []:
            if not isinstance(author, dict):
                authors.append(str(author))
                continue
            literal = author.get("name")
            if literal:
                authors.append(str(literal).strip())
            else:
                name = " ".join(
                    str(author.get(part, "")).strip()
                    for part in ("firstName", "lastName")
                    if author.get(part)
                )
                if name:
                    authors.append(name)

        extra = data.get("extra")
        pmid = data.get("PMID") or data.get("pmid")
        if not pmid and extra:
            pmid_match = re.search(r"(?:PMID|PubMed ID)\s*:\s*(\d+)", str(extra), re.I)
            pmid = pmid_match.group(1) if pmid_match else None

        return ZoteroItem(
            key=key,
            title=str(data.get("title", "") or "").strip(),
            item_type=str(data.get("itemType", "journalArticle") or "journalArticle"),
            doi=_normalise_doi(data.get("DOI") or data.get("doi")),
            pmid=_normalise_pmid(pmid),
            year=_year(data.get("date") or data.get("year")),
            journal=(data.get("publicationTitle") or data.get("journal")) or None,
            abstract=(data.get("abstractNote") or data.get("abstract")) or None,
            authors=authors,
            extra=str(extra) if extra else None,
            collection_keys=[str(v) for v in (data.get("collections") or [])],
        )

    def _attachment_from_export(
        self, data: Dict[str, Any], parent_key: str
    ) -> Optional[ZoteroAttachment]:
        key = str(data.get("key", "") or "").strip()
        if not key:
            return None
        raw_path = data.get("path") or data.get("relativePath") or data.get("filePath")
        source_path = self._resolve_export_path(raw_path, key)
        cache_path = self._find_cache(source_path, raw_path)
        return ZoteroAttachment(
            key=key,
            parent_key=parent_key,
            content_type=data.get("contentType") or data.get("mimeType"),
            link_mode=self._as_int(data.get("linkMode")),
            zotero_path=str(raw_path) if raw_path else None,
            source_path=str(source_path) if source_path else None,
            fulltext_cache_path=str(cache_path) if cache_path else None,
        )

    def _resolve_export_path(self, raw_path: Any, attachment_key: str) -> Optional[Path]:
        if not raw_path:
            return None
        value = str(raw_path).strip()
        if value.startswith("file://"):
            value = unquote(urlparse(value).path)
        if value.startswith("storage:") and self.export_json:
            value = value.split(":", 1)[1]
            candidate = self.export_json.parent / "storage" / attachment_key / value
        else:
            candidate = Path(value).expanduser()
            if not candidate.is_absolute() and self.export_json:
                candidate = self.export_json.parent / candidate
        return candidate.resolve()

    # ------------------------------ SQLite --------------------------------
    def _load_sqlite(self) -> List[ZoteroItem]:
        assert self.data_dir is not None
        db_path = self.data_dir / "zotero.sqlite"
        if not db_path.is_file():
            raise ZoteroError(f"zotero.sqlite not found in {self.data_dir}")

        uri = f"file:{db_path.as_posix()}?mode=ro"
        try:
            connection = sqlite3.connect(uri, uri=True)
            connection.row_factory = sqlite3.Row
            connection.execute("PRAGMA query_only = ON")
        except sqlite3.Error as exc:
            raise ZoteroError(f"Unable to open Zotero database read-only: {exc}") from exc

        try:
            tables = {
                row[0]
                for row in connection.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'"
                ).fetchall()
            }
            required = {"items", "itemTypes", "itemData", "fields", "itemDataValues"}
            missing = sorted(required - tables)
            if missing:
                raise ZoteroError(f"Unsupported Zotero schema; missing tables: {', '.join(missing)}")

            rows = connection.execute(
                """
                SELECT i.itemID, i.key, it.typeName
                FROM items i
                JOIN itemTypes it ON it.itemTypeID = i.itemTypeID
                WHERE lower(it.typeName) NOT IN ('attachment', 'note', 'annotation')
                ORDER BY i.itemID
                """
            ).fetchall()

            fields: Dict[int, Dict[str, str]] = {}
            for row in connection.execute(
                """
                SELECT d.itemID, f.fieldName, v.value
                FROM itemData d
                JOIN fields f ON f.fieldID = d.fieldID
                JOIN itemDataValues v ON v.valueID = d.valueID
                """
            ).fetchall():
                fields.setdefault(int(row["itemID"]), {})[str(row["fieldName"])] = row["value"]

            authors: Dict[int, List[str]] = {}
            if "itemCreators" in tables and "creators" in tables:
                for row in connection.execute(
                    """
                    SELECT ic.itemID, c.firstName, c.lastName, c.name
                    FROM itemCreators ic JOIN creators c ON c.creatorID = ic.creatorID
                    ORDER BY ic.itemID, ic.orderIndex
                    """
                ).fetchall():
                    name = row["name"] or " ".join(
                        str(row[column] or "").strip()
                        for column in ("firstName", "lastName")
                        if row[column]
                    )
                    if name:
                        authors.setdefault(int(row["itemID"]), []).append(str(name))

            collections: Dict[int, List[str]] = {}
            if "collectionItems" in tables and "collections" in tables:
                for row in connection.execute(
                    """
                    SELECT ci.itemID, c.key
                    FROM collectionItems ci JOIN collections c ON c.collectionID = ci.collectionID
                    ORDER BY ci.itemID
                    """
                ).fetchall():
                    collections.setdefault(int(row["itemID"]), []).append(str(row["key"]))

            attachment_rows: Dict[int, List[ZoteroAttachment]] = {}
            if "itemAttachments" in tables:
                for row in connection.execute(
                    """
                    SELECT ia.itemID, ia.parentItemID, ia.linkMode, ia.contentType, ia.path, ai.key
                    FROM itemAttachments ia JOIN items ai ON ai.itemID = ia.itemID
                    WHERE ia.parentItemID IS NOT NULL
                    ORDER BY ia.parentItemID, ia.itemID
                    """
                ).fetchall():
                    parent_id = int(row["parentItemID"])
                    attachment_key = str(row["key"])
                    raw_path = row["path"]
                    source_path = self._resolve_sqlite_path(
                        attachment_key, raw_path, self._as_int(row["linkMode"])
                    )
                    cache_path = self._find_cache(source_path, raw_path)
                    attachment_rows.setdefault(parent_id, []).append(
                        ZoteroAttachment(
                            key=attachment_key,
                            parent_key="",
                            content_type=row["contentType"],
                            link_mode=self._as_int(row["linkMode"]),
                            zotero_path=str(raw_path) if raw_path else None,
                            source_path=str(source_path) if source_path else None,
                            fulltext_cache_path=str(cache_path) if cache_path else None,
                        )
                    )

            items: List[ZoteroItem] = []
            for row in rows:
                item_id = int(row["itemID"])
                key = str(row["key"])
                value_map = fields.get(item_id, {})
                item = ZoteroItem(
                    key=key,
                    title=str(value_map.get("title", "") or "").strip(),
                    item_type=str(row["typeName"] or "journalArticle"),
                    doi=_normalise_doi(value_map.get("DOI") or value_map.get("doi")),
                    pmid=_normalise_pmid(value_map.get("PMID") or value_map.get("pmid")),
                    year=_year(value_map.get("date")),
                    journal=value_map.get("publicationTitle") or value_map.get("journal"),
                    abstract=value_map.get("abstractNote") or value_map.get("abstract"),
                    authors=authors.get(item_id, []),
                    extra=value_map.get("extra"),
                    collection_keys=collections.get(item_id, []),
                )
                for attachment in attachment_rows.get(item_id, []):
                    attachment.parent_key = key
                    item.attachments.append(attachment)
                items.append(item)
            return items
        except sqlite3.Error as exc:
            raise ZoteroError(f"Unable to query Zotero database: {exc}") from exc
        finally:
            connection.close()

    def _resolve_sqlite_path(
        self, attachment_key: str, raw_path: Any, link_mode: Optional[int]
    ) -> Optional[Path]:
        assert self.data_dir is not None
        value = str(raw_path or "").strip()
        if link_mode == 1:
            if value.startswith("file://"):
                value = unquote(urlparse(value).path)
            candidate = Path(value).expanduser()
            return candidate.resolve() if candidate.is_file() else candidate

        storage_dir = self.data_dir / "storage" / attachment_key
        if value.startswith("storage:"):
            value = value.split(":", 1)[1]
        if value:
            candidate = storage_dir / value
            if candidate.is_file():
                return candidate.resolve()
        if storage_dir.is_dir():
            candidates = sorted(
                p for p in storage_dir.rglob("*")
                if p.is_file() and p.name != ".zotero-ft-cache"
            )
            pdfs = [p for p in candidates if p.suffix.lower() == ".pdf"]
            if pdfs:
                return pdfs[0].resolve()
            if candidates:
                return candidates[0].resolve()
        return candidate.resolve() if value else None

    @staticmethod
    def _find_cache(source_path: Optional[Path], raw_path: Any) -> Optional[Path]:
        candidates: List[Path] = []
        if source_path:
            candidates.extend([source_path.parent / ".zotero-ft-cache", source_path.with_name(".zotero-ft-cache")])
        if raw_path:
            raw = Path(str(raw_path))
            if raw.parent != Path("."):
                candidates.append(raw.parent / ".zotero-ft-cache")
        return next((candidate.resolve() for candidate in candidates if candidate.is_file()), None)

    @staticmethod
    def _as_int(value: Any) -> Optional[int]:
        try:
            return int(value) if value is not None else None
        except (TypeError, ValueError):
            return None


class ZoteroFullTextBridge:
    """Materialize Zotero attachments and an analysis-ready JSONL corpus."""

    def __init__(self, library: ZoteroLibrary):
        self.library = library

    def materialize(
        self,
        project_dir: str,
        collection: Optional[str] = None,
        item_keys: Optional[Sequence[str]] = None,
        copy_pdfs: bool = True,
        inline_text: bool = True,
    ) -> Dict[str, Any]:
        project = Path(project_dir).expanduser().resolve()
        archive_dir = project / "original_materials" / "zotero"
        text_dir = project / "original_materials" / "fulltext"
        archive_dir.mkdir(parents=True, exist_ok=True)
        text_dir.mkdir(parents=True, exist_ok=True)

        key_filter = {str(key).casefold() for key in item_keys or []}
        items = [
            item for item in self.library.load_items(collection=collection)
            if not key_filter or item.key.casefold() in key_filter
        ]
        records: List[Dict[str, Any]] = []
        for item in items:
            records.append(
                self._materialize_item(
                    item=item,
                    archive_dir=archive_dir,
                    text_dir=text_dir,
                    copy_pdfs=copy_pdfs,
                    inline_text=inline_text,
                )
            )

        generated_at = _utc_now()
        manifest_path = project / "original_materials" / "zotero_fulltext_manifest.json"
        corpus_path = project / "original_materials" / "zotero_fulltext_corpus.jsonl"
        manifest = {
            "schema_version": 1,
            "generated_at": generated_at,
            "source": {
                "kind": self.library.source_kind,
                "path": str(self.library.data_dir or self.library.export_json),
                "collection": collection,
            },
            "item_count": len(records),
            "fulltext_count": sum(1 for record in records if record["fulltext_available"]),
            "pdf_count": sum(1 for record in records if record["pdf_available"]),
            "records": records,
        }
        _atomic_json(manifest_path, manifest)
        with corpus_path.open("w", encoding="utf-8", newline="\n") as handle:
            for record in records:
                handle.write(json.dumps(record, ensure_ascii=False, sort_keys=True) + "\n")

        return {
            "manifest_path": str(manifest_path),
            "corpus_path": str(corpus_path),
            "item_count": len(records),
            "fulltext_count": manifest["fulltext_count"],
            "pdf_count": manifest["pdf_count"],
            "missing_fulltext": [record["item_key"] for record in records if not record["fulltext_available"]],
            "records": records,
        }

    def _materialize_item(
        self,
        item: ZoteroItem,
        archive_dir: Path,
        text_dir: Path,
        copy_pdfs: bool,
        inline_text: bool,
    ) -> Dict[str, Any]:
        attachment = self._select_attachment(item.attachments)
        source_pdf = Path(attachment.source_path) if attachment and attachment.source_path else None
        cache_path = Path(attachment.fulltext_cache_path) if attachment and attachment.fulltext_cache_path else None
        archived_pdf: Optional[Path] = None
        text_path = text_dir / f"{item.key}.txt"
        text = ""
        extraction_method = "none"
        extraction_error: Optional[str] = None
        page_count: Optional[int] = None

        if copy_pdfs and source_pdf and source_pdf.is_file():
            target_name = f"{item.key}_{_slug(item.title, item.key)}{source_pdf.suffix.lower() or '.pdf'}"
            archived_pdf = archive_dir / target_name
            if not archived_pdf.exists() or _sha256(archived_pdf) != _sha256(source_pdf):
                shutil.copy2(source_pdf, archived_pdf)

        if cache_path and cache_path.is_file():
            try:
                text = cache_path.read_text(encoding="utf-8", errors="replace").strip()
                extraction_method = "zotero_ft_cache" if text else "none"
            except OSError as exc:
                extraction_error = str(exc)

        if not text and source_pdf and source_pdf.is_file():
            try:
                text, page_count = self._extract_pdf(source_pdf)
                extraction_method = "pypdf" if text else "none"
            except Exception as exc:  # keep the manifest complete for one bad PDF
                extraction_error = str(exc)

        if text:
            _atomic_text(text_path, text + "\n")

        return {
            "item_key": item.key,
            "title": item.title,
            "item_type": item.item_type,
            "doi": item.doi,
            "pmid": item.pmid,
            "year": item.year,
            "journal": item.journal,
            "authors": item.authors,
            "collection_keys": item.collection_keys,
            "attachment_key": attachment.key if attachment else None,
            "zotero_path": attachment.zotero_path if attachment else None,
            "source_pdf": str(source_pdf.resolve()) if source_pdf and source_pdf.is_file() else None,
            "archived_pdf": str(archived_pdf.resolve()) if archived_pdf and archived_pdf.is_file() else None,
            "fulltext_cache_path": str(cache_path.resolve()) if cache_path and cache_path.is_file() else None,
            "text_path": str(text_path.resolve()) if text else None,
            "pdf_available": bool(source_pdf and source_pdf.is_file()),
            "fulltext_available": bool(text),
            "extraction_method": extraction_method,
            "page_count": page_count,
            "source_pdf_sha256": _sha256(source_pdf) if source_pdf else None,
            "archived_pdf_sha256": _sha256(archived_pdf) if archived_pdf else None,
            "text_sha256": _sha256(text_path) if text else None,
            "text": text if inline_text and text else None,
            "extraction_error": extraction_error,
            "retrieved_at": _utc_now(),
        }

    @staticmethod
    def _select_attachment(attachments: Sequence[ZoteroAttachment]) -> Optional[ZoteroAttachment]:
        if not attachments:
            return None
        local = [attachment for attachment in attachments if attachment.has_local_content]
        pdfs = [
            attachment for attachment in local
            if (attachment.pdf_path is not None)
            or (attachment.content_type or "").lower() == "application/pdf"
        ]
        return (pdfs or local or list(attachments))[0]

    @staticmethod
    def _extract_pdf(path: Path) -> Tuple[str, int]:
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages: List[str] = []
            for page in reader.pages:
                pages.append(page.extract_text() or "")
            text = "\n\n".join(pages).strip()
            if text:
                return text, len(reader.pages)
        except ImportError:
            pass

        executable = shutil.which("pdftotext")
        if executable:
            with tempfile.NamedTemporaryFile(suffix=".txt", delete=False) as tmp:
                target = Path(tmp.name)
            try:
                completed = subprocess.run(
                    [executable, "-layout", str(path), str(target)],
                    capture_output=True,
                    text=True,
                    check=False,
                )
                if completed.returncode == 0 and target.is_file():
                    return target.read_text(encoding="utf-8", errors="replace").strip(), 0
            finally:
                target.unlink(missing_ok=True)
        return "", 0

    @staticmethod
    def iter_analysis_documents(corpus_path: str) -> Iterator[Dict[str, Any]]:
        """Stream records and load text from ``text_path`` when it is external."""
        path = Path(corpus_path).expanduser().resolve()
        if not path.is_file():
            raise ZoteroError(f"Zotero corpus not found: {path}")
        with path.open("r", encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                record = json.loads(line)
                if not record.get("text") and record.get("text_path"):
                    text_path = Path(record["text_path"])
                    if text_path.is_file():
                        record["text"] = text_path.read_text(encoding="utf-8", errors="replace")
                yield record

    @staticmethod
    def update_screening_workbook(
        screening_xlsx: str,
        manifest_path: str,
    ) -> Dict[str, Any]:
        """Mark Zotero-matched records as sought/retrieved without overwriting decisions."""
        try:
            import openpyxl
        except ImportError as exc:
            raise ZoteroError("openpyxl is required to update the screening workbook.") from exc

        workbook = openpyxl.load_workbook(screening_xlsx)
        sheet = workbook.active
        manifest = json.loads(Path(manifest_path).read_text(encoding="utf-8"))
        by_doi = {record["doi"].lower(): record for record in manifest.get("records", []) if record.get("doi")}
        by_pmid = {str(record["pmid"]): record for record in manifest.get("records", []) if record.get("pmid")}
        by_title = {
            re.sub(r"\W+", " ", str(record.get("title", "")).lower()).strip(): record
            for record in manifest.get("records", []) if record.get("title")
        }

        matched = 0
        retrieved = 0
        unmatched_rows: List[int] = []
        for row in range(2, sheet.max_row + 1):
            doi = _normalise_doi(sheet.cell(row=row, column=7).value)
            pmid = _normalise_pmid(sheet.cell(row=row, column=3).value)
            title = re.sub(r"\W+", " ", str(sheet.cell(row=row, column=8).value or "").lower()).strip()
            record = by_doi.get(doi or "") or by_pmid.get(pmid or "") or by_title.get(title)
            if not record:
                unmatched_rows.append(row)
                continue
            matched += 1
            sheet.cell(row=row, column=15, value="YES")
            if record.get("fulltext_available"):
                sheet.cell(row=row, column=16, value="YES")
                retrieved += 1
            note = f"Zotero item {record['item_key']}"
            existing = str(sheet.cell(row=row, column=20).value or "").strip()
            sheet.cell(row=row, column=20, value=f"{existing}; {note}".strip("; "))

        workbook.save(screening_xlsx)
        return {
            "matched_rows": matched,
            "retrieved_rows": retrieved,
            "unmatched_rows": unmatched_rows,
            "screening_xlsx": str(Path(screening_xlsx).resolve()),
        }


def _atomic_json(path: Path, value: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(_safe_json(value), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    os.replace(temp, path)


def _atomic_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(value, encoding="utf-8")
    os.replace(temp, path)
